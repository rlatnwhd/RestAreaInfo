from pathlib import Path
import sys
import json
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / '.docx-tools'))
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path.home() / 'Downloads' / '프로젝트2_기획서_김수종.docx'
OUTPUT = ROOT / 'docs' / '프로젝트2_기획서_김수종_작성완료.docx'
document = Document(SOURCE)

contents = [
    [
        '202395007',
        '김수종',
        '2026년 10월 13일',
    ],
    [
        '휴게길 — 고속도로 휴게소 통합정보',
        '고속도로 이용자는 이동 경로를 확인한 뒤에도 휴게소별 음식 가격, 입점 브랜드, 차량 편의시설을 여러 서비스에서 따로 찾아야 한다. 정보가 흩어져 있어 쉬어갈 곳을 비교하기 어렵고, 진행 방향이나 통계 기준 시점을 잘못 해석할 수 있다. 한국도로공사의 실제 공공데이터를 통합해 필요한 정보를 쉽게 확인하고, 확인되지 않은 정보는 구분해 안내하는 웹앱을 기획하였다.',
        '실제 자동차 경로와 공공데이터 기반 휴게소·음식·브랜드·주유소 시설·과거 통계를 한곳에서 확인하는 반응형 웹앱.',
    ],
    [
        '한국도로공사 고속도로 공공데이터 포털(data.ex.co.kr), 공공데이터포털(data.go.kr)의 ODCLOUD API.\n지도 및 장소·자동차 경로는 카카오 Maps / Local API와 카카오모빌리티 길찾기 API를 활용한다.',
        '제공 기관: 한국도로공사\n① 고속도로 휴게소 기준정보 현황\n② 휴게소 푸드메뉴 현황\n③ 휴게소 브랜드 매장 현황\n④ 휴게소 편의시설 현황(주유소)\n⑤ 휴게소 이용객 및 교통량 현황(15118977/v1)',
        'HTTPS GET 기반 REST API, JSON 응답을 사용한다. 한국도로공사 인증키, 공공데이터포털 인증키, 카카오 JavaScript 키 및 REST 키가 각각 필요하다. 비밀키는 Node.js 서버의 환경변수로 관리한다.\n한국도로공사: key/type 및 pageNo/numOfRows. 통계: serviceKey 및 page/perPage. 일일 호출 한도는 제공기관·신청 서비스의 승인 조건을 추가 확인하며, 서버의 10분 캐시로 반복 호출을 줄인다.',
        '기준정보: stdRestCd, unitName, routeNo, routeName.\n음식: foodNm, foodCost, etc, foodMaterial, 추천·베스트·프리미엄 및 계절 구분. 브랜드: brdName, brdDesc, stime/etime. 주유소 시설: psName, psDesc(시작·종료 구간은 영업시간으로 해석하지 않음).\n통계: 휴게소명, 노선명, 1일 이용객수(17년), 1일 차량 통행량(24년). 이용객은 2017년 명/일, 통행량은 2024년 대/일이다.\n최종수정일(lsttmAltrDttm)이 제공되는 항목에만 표시하고 조회 시각과 구분한다. 상세 갱신 주기는 추가 확인한다. 원본 X/Y 좌표계는 미검증이며 지도는 공식 주소를 카카오 주소 검색으로 변환한 위치를 사용한다.',
        '2026년 10월 13일 실제 호출 결과: 기준정보 203곳, 음식 메뉴 7,255건, 브랜드 699건, 주유소 시설 432건, 과거 통계 207곳.\n서울만남(부산)휴게소(000001)에 메뉴 54건·브랜드 2건을 공식 코드와 명칭 일치로 연결하였다. 주유소는 다른 코드(000002)여서 임의 병합하지 않았다. 주소 기반 지도 위치는 130곳을 확인하였다.\n카카오 장소 검색과 서울–부산 자동차 경로 조회도 성공하였다. 통계에는 공통 휴게소 코드가 없으므로 별도 화면에서 제공한다.',
    ],
    [
        '휴게소 정보 탐색·상세 조회 [MVP / 구현 완료]\n휴게소명·노선으로 기준정보를 검색하고, 공식 식별코드로 연결된 음식·가격·브랜드를 상세 탭에서 확인한다. 주유소 시설은 독립 목록으로 제공하며 누락·오류·연결 불일치를 구분한다.',
        '자동차 경로 및 지도 조회 [MVP / 경로·위치 표시 구현 완료]\n출발지와 도착지의 주소·장소 검색 결과를 선택하면 실제 자동차 경로와 API가 제공하는 거리·소요시간을 표시한다. 공식 주소에서 확인한 휴게소 위치를 지도에 표시한다.\n단, 진입 방향·접속 도로가 미확인되어 경로상 휴게소 판별과 주행 순서 정렬은 후속 검증 과제로 남긴다.',
        '휴게소 비교 및 과거 통계 시각화 [MVP / 구현 완료]\n휴게소 2~3곳의 메뉴·가격·브랜드·연결 가능한 주유소 시설을 비교한다. 2017년 이용객과 2024년 차량 통행량은 연도·단위를 구분해 독립 차트로 제공한다. 현재 혼잡도나 대기시간으로 표현하지 않는다.',
        '① 사이트 접속 → ② 출발지·도착지 검색 후 장소 선택 → ③ 실제 자동차 경로 확인 → ④ 전체 휴게소 목록에서 이름·노선 검색 → ⑤ 휴게소 상세 메뉴·브랜드 확인 → ⑥ 2~3곳 선택 후 비교 → ⑦ 과거 통계 및 출처 확인.\n현재 목록은 전체 휴게소이며, 경로상 진입 가능한 휴게소로 확정된 결과가 아님을 명확히 안내한다.',
    ],
    [
        '메인(경로 찾기): 데이터 요약, 출발·도착 검색, 실제 지도와 전체 휴게소 목록.\n휴게소 찾기: 이름·노선 필터, 상세 탭(기본정보/음식/브랜드/주유소 시설/과거 통계), 비교 선택.\n음식 메뉴: 메뉴·휴게소 검색과 등록 가격 표. 주유소 시설: 독립 시설 목록.\n과거 통계: 연도별 차트와 원본 수치. 이용·데이터 안내: 출처·조회 시각·연결 방식·한계. 비교 화면은 모달로 구성한다.',
        '사용한 도구: ChatGPT의 Codex. 기존 프로젝트 코드 작성·수정과 반응형 화면 제작에 활용하였다.\n작업 요청을 정리한 프롬프트: “React + Vite로 고속도로 휴게소 공공데이터 웹사이트를 만들고, 공공서비스처럼 검색·지도·정보가 명확히 구분되는 한국어 UI로 디자인해 줘. 가로폭이 줄면 자연스럽게 정돈되도록 반응형으로 만들고 실제 데이터, 출처 및 누락 상태를 표시해 줘.”\n아래 이미지는 AI로 제작한 UI의 실제 구현 화면이다.',
    ],
    [
        '지도: 카카오 지도에 실제 자동차 경로와 공식 주소에서 확인한 위치 표시. 목록: 휴게소명·노선·정보 상태와 비교 선택.\n표/카드: 음식 가격, 입점 브랜드, 주유소 시설. 차트: Recharts 막대그래프로 2017년 이용객 또는 2024년 통행량을 분리 표시.\n모바일에서는 검색창·지도·목록을 세로로 재배치하고, 표를 읽기 쉬운 카드 형태로 전환한다. 로딩·오류·정보 없음과 키보드 초점을 처리한다.',
    ],
    [
        '프론트엔드: React + Vite, JavaScript/JSX, 일반 CSS. 컴포넌트 재사용과 빠른 개발을 위해 선택하였다. 지도는 Kakao Maps, 통계는 Recharts를 사용한다.\n백엔드: Node.js HTTP 서버와 Fetch API. 비밀키 보호, 외부 API 호출, 페이지 수집·정규화·공식 코드 검증 및 메모리 캐시를 담당한다. 별도 DB·로그인 기능은 MVP에 포함하지 않는다.',
        '배포 계획: Render의 Node.js Web Service에 프론트 빌드와 프록시 서버를 함께 배포할 예정이다. 정적 호스팅만으로는 서버 전용 인증키를 안전하게 처리할 수 없기 때문이다.\n배포 전 PORT/0.0.0.0 호스트 바인딩, 서버 환경변수, 카카오 지도 배포 도메인 등록, HTTPS 및 API 이용 한도를 확인한다. 빌드 후 npm start로 운영한다. 현재는 로컬 구현·검증 단계이며 공개 배포는 아직 하지 않았다.\n목표: 10월 29일까지 배포 및 최종 점검 완료. 참고: render.com/docs/web-services',
    ],
    [
        '휴게소 음식 가격 비교 중심의 유사 서비스에서 확장하여 실제 자동차 경로, 기준정보, 음식·브랜드·주유소 시설·과거 통계를 통합 탐색한다. 공식 식별코드로 연결하고, 진입 방향 미검증·주유소 코드 차이·통계 기준 연도를 명확히 안내하는 데이터 정확성을 차별점으로 삼는다.',
        '이동 전 휴게소 정보를 한곳에서 확인해 검색 시간을 줄이고, 메뉴 가격과 브랜드를 비교해 쉬어갈 곳을 선택하는 데 도움을 준다. PC·모바일에서 같은 정보를 쉽게 이용하며, 실제 공공데이터의 수집·결합·시각화 사례를 구현한다. 정보의 출처와 한계를 함께 제시해 잘못된 판단을 줄인다.',
    ],
]

def fill_cell(cell, content):
    cell.text = ''
    for index, line in enumerate(content.split('\n')):
        paragraph = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.line_spacing = 1.13
        run = paragraph.add_run(line)
        run.font.name = '맑은 고딕'
        run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '맑은 고딕')
        run.font.size = Pt(10)

assert len(document.tables) == 9
for table_index, values in enumerate(contents):
    table = document.tables[table_index]
    assert len(table.rows) == len(values)
    for row, content in zip(table.rows, values):
        fill_cell(row.cells[1], content)
        tr_pr = row._tr.get_or_add_trPr()
        if tr_pr.find(qn('w:cantSplit')) is None:
            tr_pr.append(OxmlElement('w:cantSplit'))
        for paragraph in row.cells[0].paragraphs:
            paragraph.paragraph_format.space_after = Pt(4)
            for run in paragraph.runs:
                run.font.name = '맑은 고딕'
                run.font.size = Pt(10)
        row.height = None

schedule = [
    '기획서 브리핑: 주제·문제 정의·공공 API 활용 방식을 설명하고 실제 호출 결과와 화면 시안을 제시한다. 구현 완료 범위와 진입 방향·통계 연결 등 남은 검증 과제를 정리한다.',
    '개발 점검: 데이터 연결·상세·비교·통계 화면을 보완한다. 방향·접속 도로의 추가 공식 자료를 조사하고, 가격 누락·코드 불일치·모바일 UI를 검증한다. 배포 환경 설정과 카카오 도메인 등록을 준비한다.',
    '발표: 10월 29일까지 공개 배포와 최종 점검을 완료하고 배포 URL에서 기능을 시연한다. 공공데이터 결합 방식, 검증 결과, 남은 제한 및 생성형 AI 활용 과정·직접 검증한 내용을 발표 자료에 설명한다.',
]
for row, content in zip(document.tables[8].rows[1:], schedule):
    fill_cell(row.cells[1], content)
    row.height = None

for paragraph in document.paragraphs:
    if paragraph.text.strip().startswith(tuple(f'{i}.' for i in range(1,8))):
        paragraph.paragraph_format.keep_with_next = True
    if '와이어프레임 / 스케치 삽입란' in paragraph.text:
        paragraph.paragraph_format.page_break_before = True
        paragraph.paragraph_format.keep_with_next = True
    if '(AI로 제작한 화면 스케치 이미지를 삽입하세요)' in paragraph.text:
        paragraph.clear()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        picture = ROOT / 'artifacts' / 'plan-home.png'
        if picture.exists():
            paragraph.add_run().add_picture(str(picture), width=Cm(16))
        else:
            paragraph.add_run().add_picture(str(ROOT / 'artifacts' / 'desktop.png'), width=Cm(14))
        caption = paragraph.insert_paragraph_before('휴게길 메인 화면 — 실제 구현 화면(2026.10.13). 지도·조회 결과는 API 연결 및 조회 시점에 따라 달라질 수 있음.')
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.paragraph_format.keep_with_next = True
        for run in caption.runs:
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor.from_string('64748B')

for section in document.sections:
    footer = section.footer.paragraphs[0]
    if not footer.text.strip():
        footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = footer.add_run('휴게길 · 프로젝트 2 기획서 | 김수종 | ')
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string('64748B')
        field = OxmlElement('w:fldSimple')
        field.set(qn('w:instr'), 'PAGE')
        footer._p.append(field)

document.core_properties.title = '휴게길 — 프로젝트 2 기획서'
document.core_properties.author = '김수종'
document.core_properties.subject = '공공데이터 활용 고속도로 휴게소 통합정보 웹앱'
OUTPUT.parent.mkdir(exist_ok=True)
document.save(OUTPUT)

checked = Document(OUTPUT)
assert len(checked.tables) == 9
assert len(checked.inline_shapes) >= 1
for table in checked.tables:
    for row in table.rows:
        assert row.cells[1].text.strip()
assert checked.tables[0].cell(0,1).text == '202395007'
assert checked.tables[0].cell(2,1).text == '2026년 10월 13일'
print(json.dumps({'saved': str(OUTPUT), 'tables':len(checked.tables), 'images':len(checked.inline_shapes)}, ensure_ascii=True))
